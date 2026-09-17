import pytest
from sqlalchemy import create_engine
import pymysql
import oracledb
from commonUtilities.utility import BaseUtility
from testConfigurations.config import *
baseUtility = BaseUtility()


@pytest.fixture(scope="module")
def connect_to_mysqldb():
    try:
        baseUtility.log_info("connecting to mysql server...")
        mysql_conn = create_engine(F"mysql+pymysql://{MYSQL_USERNAME}:{MYSQL_PASSWORD}@{MYSQL_HOST}:{MYSQL_PORT}/{MYSQL_DATABASE}")
        connection=mysql_conn.connect()
        baseUtility.log_info("mysql server connected.")
        yield connection
        connection.close()
    except Exception as e:
        baseUtility.log_error(f"failed to connect mysql server with exception : {e}")

@pytest.fixture(scope="module")
def connect_to_oracledb():
    try:
        baseUtility.log_info("connecting to oracle server...")
        oracle_conn = create_engine(f"oracle+oracledb://{ORACLE_USERNAME}:{ORACLE_PASSWORD}@{ORACLE_HOST}:{ORACLE_PORT}/{ORACLE_SERVICE}")
        connection=oracle_conn.connect()
        baseUtility.log_info("oracle server connected.")
        yield connection
        connection.close()
    except Exception as e:
        baseUtility.log_error(f"failed to connect oracle server with exception : {e}")



import os
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

REPORT_FILE = "etl_test_results.xlsx"

PASS_FILL = PatternFill(
    start_color="90EE90",
    end_color="90EE90",
    fill_type="solid"
)

FAIL_FILL = PatternFill(
    start_color="FF0000",
    end_color="FF0000",
    fill_type="solid"
)


def initialize_report():
    if not os.path.exists(REPORT_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Results"
        ws.append(["Timestamp", "Test Case", "Status"])
        wb.save(REPORT_FILE)


@pytest.hookimpl(hookwrapper=True)
def pytest_runtest_makereport(item, call):
    outcome = yield
    report = outcome.get_result()

    if report.when == "call":
        initialize_report()

        wb = load_workbook(REPORT_FILE)
        ws = wb["Test Results"]

        status = "PASS" if report.passed else "FAIL"

        ws.append([
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            item.name,
            status
        ])

        row_num = ws.max_row

        if report.failed:
            for cell in ws[row_num]:
                cell.fill = FAIL_FILL
        else:
            for cell in ws[row_num]:
                cell.fill = PASS_FILL

        wb.save(REPORT_FILE)

'''
ws.append([
    datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    item.name,
    source_table,
    target_table,
    validation_type,
    status
])
'''


